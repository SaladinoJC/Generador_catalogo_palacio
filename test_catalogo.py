"""Regresiones del listado final y de los productos personalizados.

Ejecutar: python -m unittest -v
"""
import contextlib
import io
import re
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from PIL import Image

import generar_catalogo as catalogo
import sincronizar_fotos as fotos


class CatalogoFinalTests(unittest.TestCase):
    def test_seleccion_y_orden_del_listado_con_personalizado_al_final(self):
        wb = load_workbook(catalogo.BASE / "PRODUCTOS CATALOGO NUEVO.xlsx", data_only=True)
        listado = [r for r in wb.worksheets[0].values if isinstance(r[1], int)]
        wb.close()
        productos = catalogo.load_products()
        self.assertEqual(len(productos), 101)
        self.assertEqual([p["sku"] for p in productos[:-1]], [str(r[2]) for r in listado[:-1] if str(r[2]) != "1099"])
        self.assertEqual(next(p["price"] for p in productos if p["sku"] == "122"), 52000)
        self.assertFalse(any(p["custom"] for p in productos[:-1]))
        personalizado = productos[-1]
        self.assertEqual(personalizado["sku"], "")
        self.assertTrue(personalizado["custom"])
        self.assertEqual(personalizado["img"].name, "personalizado.jpg")
        for condicion in ("medidas", "con borde o sin borde", "vectorizado", "Corel (.cdr)", "SVG (.svg)", "50%", "seña"):
            self.assertIn(condicion, personalizado["desc"])

    def test_todas_las_fotos_abren_y_los_enlaces_ausentes_van_al_inicio(self):
        productos = catalogo.load_products()
        for producto in productos:
            with self.subTest(producto=producto["name"]):
                with Image.open(producto["img"]) as imagen:
                    imagen.verify()
        sin_ficha = {"5400", "3807", "122", ""}
        self.assertEqual({p["sku"] for p in productos if p["url"] == catalogo.WEB_URL}, sin_ficha)
        for valor in (None, "", " ", "/"):
            self.assertEqual(catalogo.build_product_url(valor), catalogo.WEB_URL)

    def test_paginacion_preserva_orden_y_separa_personalizados_en_ambos_formatos(self):
        productos = catalogo.load_products()
        for cantidad in (3, 6):
            with self.subTest(productos_por_pagina=cantidad):
                paginas = catalogo.product_batches(productos, cantidad)
                self.assertEqual([p for pagina in paginas for p in pagina], productos)
                self.assertEqual(paginas[-1], [productos[-1]])
                self.assertTrue(all(len(pagina) <= cantidad for pagina in paginas))


class SincronizacionTests(unittest.TestCase):
    def test_sincronizar_no_duplica_la_foto_ni_sobrescribe_la_fila_sin_sku(self):
        with tempfile.TemporaryDirectory() as directorio:
            base = Path(directorio)
            imagenes = base / "imagenes"
            imagenes.mkdir()
            for nombre in ("personalizado.jpg", "123.jpg"):
                Image.new("RGB", (4, 4), "white").save(imagenes / nombre)
            excel = base / "productos.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Productos"
            ws.append(["SKU", "Producto", "Descripción", "Precio", "Imagen", "Mostrar", "URL", "Tipo"])
            ws.append([None, "Felpudo personalizado", "A medida", None, "imagenes/personalizado.jpg", "SI", None, "Personalizado"])
            wb.save(excel)
            wb.close()
            with patch.object(fotos, "EXCEL", excel), patch.object(fotos, "IMG_DIR", imagenes), patch("builtins.input", return_value="n"), contextlib.redirect_stdout(io.StringIO()):
                fotos.sincronizar()
            wb = load_workbook(excel)
            ws = wb["Productos"]
            self.assertEqual(ws.max_row, 3)
            self.assertIsNone(ws.cell(2, 1).value)
            self.assertEqual(ws.cell(2, 2).value, "Felpudo personalizado")
            self.assertEqual(ws.cell(2, 8).value, "Personalizado")
            self.assertEqual(ws.cell(3, 1).value, "123")
            self.assertIsNone(ws.cell(3, 7).value)
            wb.close()


class VersionImpresionTests(unittest.TestCase):
    def test_exporta_aparte_sin_precios_ofertas_botones_ni_enlaces(self):
        productos = catalogo.load_products()
        oferta = dict(productos[0], price=52000, old_price=65000, offer=True)
        canvas_original = catalogo.canvas.Canvas

        def canvas_sin_comprimir(*args, **kwargs):
            return canvas_original(*args, pageCompression=0, **kwargs)

        def sin_imagenes(pdf):
            # Las imágenes codificadas pueden contener palabras por azar.
            objetos = re.findall(rb"(?m)^\d+ \d+ obj\n(.*?)\nendobj", pdf, re.DOTALL)
            return b"\n".join(objeto for objeto in objetos if b"/Subtype /Image" not in objeto)

        with tempfile.TemporaryDirectory() as directorio:
            digital = Path(directorio) / "digital.pdf"
            impresion = Path(directorio) / "impresion.pdf"
            digital.write_bytes(b"version digital existente")
            for columnas in (1, 2):
                with self.subTest(columnas=columnas), patch.object(catalogo, "COLUMNAS", columnas), patch.object(catalogo, "SALIDA", digital), patch.object(catalogo, "SALIDA_IMPRESION", impresion), patch.object(catalogo, "load_products", return_value=[oferta, productos[-1]]), patch.object(catalogo.canvas, "Canvas", side_effect=canvas_sin_comprimir), contextlib.redirect_stdout(io.StringIO()):
                    anterior = digital.read_bytes()
                    catalogo.generate(print_version=True)
                    self.assertEqual(digital.read_bytes(), anterior)
                    archivo_impresion = impresion.read_bytes()
                    contenido = sin_imagenes(archivo_impresion)
                    for prohibido in (b"52.000", b"65.000", b"PRECIO", b"Precios", b"OFERTA", b"Ver en la web", b"TOCA", b"Cotizar por WhatsApp", b"/Subtype /Link"):
                        self.assertFalse(prohibido in contenido, prohibido.decode())
                    self.assertIn(b"50%", contenido)
                    self.assertIn(b"Felpudos personalizados", contenido)
                    catalogo.generate()
                    contenido_digital = sin_imagenes(digital.read_bytes())
                    for requerido in (b"52.000", b"65.000", b"OFERTA", b"Ver en la web", b"/Subtype /Link"):
                        self.assertIn(requerido, contenido_digital)
                    self.assertEqual(impresion.read_bytes(), archivo_impresion)


if __name__ == "__main__":
    unittest.main()
