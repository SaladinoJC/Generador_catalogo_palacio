from pathlib import Path
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor, white
from reportlab.lib.utils import ImageReader
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph
from xml.sax.saxutils import escape
from PIL import Image
from datetime import datetime
import textwrap
import urllib.parse
import io
import sys
import argparse

# ==============================================================================
# CONFIGURACIÓN DEL COMERCIO Y CONTACTO
# ==============================================================================
NOMBRE_COMERCIO = "El Palacio de la Limpieza"
SUCURSALES = "Moreno 3260  |  Av. Edison 333"
HORARIOS = "Lun a Vie 8 a 18 hs  ·  Sáb 8 a 13 hs"
WEB_URL = "https://www.elpalaciodelalimpieza.com.ar"
WEB_DISPLAY = "www.elpalaciodelalimpieza.com.ar"
INSTAGRAM_USER = "@palaciodelalimpieza"
INSTAGRAM_URL = "https://www.instagram.com/palaciodelalimpieza"
WHATSAPP_NUMERO = "5492234173945"
WHATSAPP_DISPLAY = "(223) 417-3945"
WHATSAPP_TEXTO_GENERAL = "¡Hola! Quisiera hacer una consulta sobre el catálogo de productos."

# ==============================================================================
# CONFIGURACIÓN DE DISEÑO DE PÁGINA
# 2 = Cuadrícula de 2 columnas (6 productos por página, 2 columnas x 3 filas)
# 1 = Modo lista horizontal (3 productos por página)
# ==============================================================================
COLUMNAS = 2
INCLUIR_PORTADA = True
TITULO_PORTADA = "CATÁLOGO DE ARTÍCULOS"
DISCLAIMER_PORTADA = "Precios del día sujetos a cambios o modificaciones"

# ==============================================================================
# RUTAS Y COLORES INSTITUCIONALES (Verde estandarizado con el fondo del logo)
# ==============================================================================
BASE = Path(__file__).resolve().parent
EXCEL = BASE / "catalogo_productos.xlsx"
LOGO = BASE / "logo.jpg"
PORTADIN_IMG = BASE / "portadin.png"
CAMIONETA_IMG = BASE / "camioneta.jpg"
PORTADA_CACHE = BASE / "salida" / "_portada_compuesta.png"
SALIDA = BASE / "salida" / "catalogo.pdf"
SALIDA_IMPRESION = BASE / "salida" / "catalogo_impresion.pdf"

# Verde extraído directamente de los píxeles del logo.jpg para coincidencia 100% exacta
GREEN = HexColor("#146B41")
DARK = HexColor("#0F3E28")
LIGHT = HexColor("#EEF5F1")
MID = HexColor("#556B61")
LINE = HexColor("#DDEDE5")
RED = HexColor("#DC2626")

W, H = A4


def link_url_new_tab(c, url, rect, print_version=False):
    """Enlace estándar de ReportLab sin bordes para compatibilidad universal."""
    if not print_version:
        c.linkURL(url, rect, thickness=0)


def money(value):
    try:
        n = float(value)
    except (TypeError, ValueError):
        return str(value or "")
    return "$ " + f"{n:,.0f}".replace(",", ".")


def get_mes_actual():
    meses = {
        1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
        7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre"
    }
    now = datetime.now()
    return f"{meses.get(now.month, '')} {now.year}"


def build_product_url(val):
    if not val:
        return WEB_URL
    val = str(val).strip()
    if not val:
        return WEB_URL
    if val.startswith("http://") or val.startswith("https://"):
        return val
    if val.startswith("www."):
        return f"https://{val}"
    for prefix in ("elpalaciodelalimpieza.com.ar/", "elpalaciodelalimpieza.com.ar"):
        if val.startswith(prefix):
            val = val[len(prefix):]
    slug = val.lstrip("/")
    if not slug:
        return WEB_URL
    return f"https://www.elpalaciodelalimpieza.com.ar/{slug}"


def load_products():
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb["Productos"]

    headers = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if v:
            headers[str(v).strip()] = col

    required = ["SKU", "Producto", "Descripción", "Precio", "Mostrar"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise RuntimeError("Faltan columnas obligatorias en el Excel: " + ", ".join(missing))

    products = []
    for row in range(2, ws.max_row + 1):
        mostrar = str(ws.cell(row, headers["Mostrar"]).value or "").strip().upper()
        if mostrar not in ("SI", "SÍ", "YES", "1", "TRUE"):
            continue

        sku = str(ws.cell(row, headers["SKU"]).value or "").strip()
        nombre = str(ws.cell(row, headers["Producto"]).value or "").strip()
        if not nombre:
            continue

        personalizado = (
            "Tipo" in headers
            and str(ws.cell(row, headers["Tipo"]).value or "").strip().casefold() == "personalizado"
        )

        imagen_val = ""
        if "Imagen" in headers:
            imagen_val = str(ws.cell(row, headers["Imagen"]).value or "").strip()

        imagen_path = None
        if imagen_val:
            candidate = BASE / imagen_val
            if candidate.exists():
                imagen_path = candidate
            else:
                candidate_sub = BASE / "imagenes" / imagen_val
                if candidate_sub.exists():
                    imagen_path = candidate_sub

        if sku and (not imagen_path or not imagen_path.exists()):
            for ext in (".png", ".jpg", ".jpeg", ".webp", ".PNG", ".JPG"):
                candidate_sku = BASE / "imagenes" / f"{sku}{ext}"
                if candidate_sku.exists():
                    imagen_path = candidate_sku
                    break
            if not imagen_path:
                imagen_path = BASE / "imagenes" / f"{sku}.png"

        presentacion = ""
        if "Presentación" in headers:
            presentacion = str(ws.cell(row, headers["Presentación"]).value or "").strip()

        fragancias = ""
        if "Fragancias" in headers:
            fragancias = str(ws.cell(row, headers["Fragancias"]).value or "").strip()

        url_raw = ""
        if "URL" in headers:
            url_raw = str(ws.cell(row, headers["URL"]).value or "").strip()
        elif "Url" in headers:
            url_raw = str(ws.cell(row, headers["Url"]).value or "").strip()

        oferta = False
        if "Oferta" in headers:
            v_of = str(ws.cell(row, headers["Oferta"]).value or "").strip().upper()
            oferta = v_of in ("SI", "SÍ", "YES", "1", "TRUE", "OFERTA")
        elif "Destacado" in headers:
            v_of = str(ws.cell(row, headers["Destacado"]).value or "").strip().upper()
            oferta = v_of in ("SI", "SÍ", "YES", "1", "TRUE", "DESTACADO")

        precio_anterior = None
        for h_cand in ("PrecioAnterior", "Precio Anterior", "Precio_Anterior", "Precio anterior"):
            if h_cand in headers:
                raw_ant = ws.cell(row, headers[h_cand]).value
                if raw_ant is not None and str(raw_ant).strip() != "":
                    try:
                        precio_anterior = float(raw_ant)
                    except (ValueError, TypeError):
                        precio_anterior = None
                break

        categoria = ""
        if "Categoría" in headers:
            categoria = str(ws.cell(row, headers["Categoría"]).value or "").strip()
        elif "Categoria" in headers:
            categoria = str(ws.cell(row, headers["Categoria"]).value or "").strip()

        products.append({
            "sku": sku,
            "name": nombre,
            "desc": str(ws.cell(row, headers["Descripción"]).value or "").strip(),
            "price": ws.cell(row, headers["Precio"]).value,
            "old_price": precio_anterior,
            "img": imagen_path,
            "presentation": presentacion,
            "fragrances": fragancias,
            "url": build_product_url(url_raw),
            "offer": oferta,
            "category": categoria,
            "custom": personalizado,
        })

    wb.close()
    return products


def draw_image_placeholder(c, x, y, w, h):
    """Dibuja el recuadro con icono de imagen estilizado de diseño.png cuando no hay foto."""
    c.setFillColor(LIGHT)
    c.roundRect(x, y, w, h, 8, fill=1, stroke=0)

    icon_w = 34
    icon_h = 26
    ix = x + (w - icon_w) / 2
    iy = y + (h - icon_h) / 2 + 7

    c.setStrokeColor(MID)
    c.setLineWidth(1.3)
    c.roundRect(ix, iy, icon_w, icon_h, 3, fill=0, stroke=1)

    c.circle(ix + 9, iy + 17, 2.5, fill=0, stroke=1)

    p = c.beginPath()
    p.moveTo(ix + 3, iy + 4)
    p.lineTo(ix + 13, iy + 15)
    p.lineTo(ix + 19, iy + 9)
    p.lineTo(ix + 25, iy + 14)
    p.lineTo(ix + 31, iy + 4)
    c.drawPath(p, stroke=1, fill=0)

    c.setFillColor(MID)
    c.setFont("Helvetica-Bold", 6.8)
    c.drawCentredString(x + w / 2, y + (h - icon_h) / 2 - 8, "IMAGEN NO ENCONTRADA")


def fit_image(c, path, x, y, box_w, box_h):
    if not path or not path.exists():
        draw_image_placeholder(c, x, y, box_w, box_h)
        return

    try:
        c.setFillColor(white)
        c.roundRect(x, y, box_w, box_h, 8, fill=1, stroke=0)

        with Image.open(path) as im:
            iw, ih = im.size
            scale = min(box_w / iw, box_h / ih)
            w, h = iw * scale, ih * scale

            target_px_w = max(1, int(w * 200 / 72))
            target_px_h = max(1, int(h * 200 / 72))

            if iw > target_px_w or ih > target_px_h:
                im_opt = im.resize((target_px_w, target_px_h), Image.Resampling.LANCZOS)
            else:
                im_opt = im.copy()

            if im_opt.mode in ("RGBA", "LA") or (im_opt.mode == "P" and "transparency" in im_opt.info):
                im_rgba = im_opt.convert("RGBA")
                bg = Image.new("RGB", im_rgba.size, (255, 255, 255))
                bg.paste(im_rgba, mask=im_rgba.split()[3])
                final_img = bg
            else:
                final_img = im_opt.convert("RGB")

            buf = io.BytesIO()
            final_img.save(buf, format="JPEG", quality=85, optimize=True)
            buf.seek(0)
            img_reader = ImageReader(buf)

        c.drawImage(
            img_reader,
            x + (box_w - w) / 2,
            y + (box_h - h) / 2,
            width=w,
            height=h,
            preserveAspectRatio=True,
        )
    except Exception:
        draw_image_placeholder(c, x, y, box_w, box_h)


def wrapped_lines(text, width_chars, max_lines):
    lines = textwrap.wrap(text or "", width=width_chars)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        if lines:
            lines[-1] = lines[-1].rstrip(" .") + "..."
    return lines


def draw_diagonal_offer_ribbon(c, card_x, card_y, card_w, card_h, offset=24, band_w=16, font_sz=7):
    """Dibuja una cinta diagonal de ¡OFERTA! en la esquina superior derecha."""
    c.saveState()
    clip_path = c.beginPath()
    clip_path.roundRect(card_x, card_y, card_w, card_h, 8)
    c.clipPath(clip_path, stroke=0, fill=0)

    cx = card_x + card_w
    cy = card_y + card_h

    # Centro de la cinta diagonal (-45 grados: de arriba-izquierda a abajo-derecha)
    c.translate(cx - offset, cy - offset)
    c.rotate(-45)

    c.setFillColor(RED)
    c.rect(-50, -band_w / 2, 100, band_w, fill=1, stroke=0)

    c.setFillColor(white)
    c.setFont("Helvetica-Bold", font_sz)
    c.drawCentredString(0, -font_sz * 0.35, "¡OFERTA!")

    c.restoreState()


def draw_header(c, page_no, total_pages=1, print_version=False):
    header_h = 96
    c.setFillColor(GREEN)
    c.rect(0, H - header_h, W, header_h, fill=1, stroke=0)

    # Logo institucional
    if LOGO.exists():
        c.drawImage(str(LOGO), 20, H - header_h + 15, width=66, height=66,
                    preserveAspectRatio=True, mask="auto")

    # Título y datos principales
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 19)
    c.drawString(98, H - 32, "CATÁLOGO DE PRODUCTOS")

    c.setFont("Helvetica-Bold", 9.5)
    c.drawString(99, H - 47, f"{NOMBRE_COMERCIO} · Mar del Plata")

    c.setFont("Helvetica", 7.8)
    c.drawString(99, H - 61, f"Sucursales: {SUCURSALES}")
    c.drawString(99, H - 74, f"Horarios: {HORARIOS}")

    # Lado derecho: Paginación y vigencia
    pag_text = f"Página {page_no} de {total_pages}" if total_pages > 1 else f"Página {page_no}"
    c.setFont("Helvetica-Bold", 8.5)
    c.drawRightString(W - 20, H - 32, pag_text)

    c.setFont("Helvetica", 7.8)
    vigencia = "Catálogo para consulta en el local" if print_version else f"Precios vigentes: {get_mes_actual()}"
    c.drawRightString(W - 20, H - 46, vigencia)

    # Botón de WhatsApp muy visible y llamativo en la cabecera
    if WHATSAPP_NUMERO:
        if print_version:
            c.setFont("Helvetica-Bold", 9)
            c.drawRightString(W - 20, H - 74, f"WhatsApp: {WHATSAPP_DISPLAY}")
            return
        wa_h_w = 175
        wa_h_h = 24
        wa_h_x = W - 20 - wa_h_w
        wa_h_y = H - 80

        c.setFillColor(white)
        c.roundRect(wa_h_x, wa_h_y, wa_h_w, wa_h_h, 6, fill=1, stroke=0)

        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 7.8)
        c.drawCentredString(wa_h_x + wa_h_w / 2, wa_h_y + 7.5, f"TOCA AQUÍ: WhatsApp {WHATSAPP_DISPLAY}")

        wa_top_url = f"https://wa.me/{WHATSAPP_NUMERO}?text={urllib.parse.quote(WHATSAPP_TEXTO_GENERAL)}"
        link_url_new_tab(c, wa_top_url, (wa_h_x, wa_h_y, wa_h_x + wa_h_w, wa_h_y + wa_h_h))


def draw_footer(c, print_version=False):
    footer_h = 26
    c.setFillColor(GREEN)
    c.rect(0, 0, W, footer_h, fill=1, stroke=0)
    c.setFillColor(white)

    # Enlace a Página Web
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(20, 9, WEB_DISPLAY)
    web_w = c.stringWidth(WEB_DISPLAY, "Helvetica-Bold", 7.5)
    link_url_new_tab(c, WEB_URL, (20, 2, 20 + web_w, footer_h), print_version)

    # Separador
    sep_x = 20 + web_w + 8
    c.setFont("Helvetica", 7.5)
    c.drawString(sep_x, 9, "|")

    # Enlace a Instagram
    ig_x = sep_x + 8
    c.setFont("Helvetica-Bold", 7.5)
    ig_label = f"Instagram: {INSTAGRAM_USER}"
    c.drawString(ig_x, 9, ig_label)
    ig_w = c.stringWidth(ig_label, "Helvetica-Bold", 7.5)
    link_url_new_tab(c, INSTAGRAM_URL, (ig_x, 2, ig_x + ig_w, footer_h), print_version)

    # Enlace a Pedidos por WhatsApp
    if WHATSAPP_NUMERO:
        c.setFont("Helvetica-Bold", 7.5)
        wa_display = f"WhatsApp: {WHATSAPP_DISPLAY}" if print_version else f"Toca para WhatsApp: {WHATSAPP_DISPLAY}"
        c.drawRightString(W - 20, 9, wa_display)
        wa_w = c.stringWidth(wa_display, "Helvetica-Bold", 7.5)
        wa_url = f"https://wa.me/{WHATSAPP_NUMERO}?text={urllib.parse.quote(WHATSAPP_TEXTO_GENERAL)}"
        link_url_new_tab(c, wa_url, (W - 20 - wa_w, 2, W - 20, footer_h), print_version)
    else:
        c.setFont("Helvetica", 7.5)
        c.drawRightString(W - 20, 9, "Consultanos en el local" if print_version else "Precios expresados en pesos argentinos")


# ==============================================================================
# DISEÑO DE TARJETA EN CUADRÍCULA (RÉPLICA EXACTA DE diseño.png)
# 3 zonas horizontales: Barra SKU/Tag, Cuerpo con Foto Cuadrada + Título, Pie con Precio + Botón Web
# ==============================================================================
def draw_card_grid(c, p, x, y, w, h, print_version=False):
    # Contenedor de la tarjeta (fondo blanco y borde redondeado)
    c.setFillColor(white)
    c.setStrokeColor(LINE)
    c.setLineWidth(0.9)
    c.roundRect(x, y, w, h, 8, fill=1, stroke=1)

    # --------------------------------------------------------------------------
    # ZONA 1: BARRA SUPERIOR (SKU + BADGE DE PRESENTACIÓN/FRAGANCIAS EXTENDIDO)
    # --------------------------------------------------------------------------
    header_bar_y = y + h - 28
    c.setStrokeColor(LINE)
    c.setLineWidth(0.8)
    c.line(x, header_bar_y, x + w, header_bar_y)

    # SKU
    c.setFillColor(DARK)
    c.setFont("Helvetica-Bold", 8.5)
    sku_str = f"SKU {p['sku']}" if p["sku"] else "A PEDIDO"
    c.drawString(x + 12, y + h - 18, sku_str)
    sku_w = c.stringWidth(sku_str, "Helvetica-Bold", 8.5)

    # Límite derecho para el badge de presentación/fragancias
    if p.get("offer") and not print_version:
        max_right = x + w - 44
    else:
        max_right = x + w - 12

    # El badge ahora puede ocupar todo el ancho disponible hasta el SKU
    min_left = x + 12 + sku_w + 8
    avail_tag_w = max(40, max_right - min_left)

    tags = []
    if p["presentation"]:
        tags.append(p["presentation"])
    if p["fragrances"]:
        tags.append(p["fragrances"])

    if tags:
        tag_text = " · ".join(tags)
        c.setFont("Helvetica-Bold", 7.2)

        # Truncar suavemente solo si excede el renglón completo hasta el SKU
        while c.stringWidth(tag_text, "Helvetica-Bold", 7.2) + 14 > avail_tag_w and len(tag_text) > 4:
            tag_text = tag_text[:-1]
        if len(tag_text) < len(" · ".join(tags)):
            tag_text = tag_text.rstrip(" ·,.-") + "..."

        tag_w = min(avail_tag_w, c.stringWidth(tag_text, "Helvetica-Bold", 7.2) + 14)
        tag_x = max_right - tag_w

        c.setFillColor(LIGHT)
        c.roundRect(tag_x, y + h - 21.5, tag_w, 15, 5, fill=1, stroke=0)
        c.setFillColor(DARK)
        c.setFont("Helvetica-Bold", 7.2)
        c.drawCentredString(tag_x + tag_w / 2, y + h - 17.5, tag_text)

    # Cinta diagonal de ¡OFERTA!
    if p.get("offer") and not print_version:
        draw_diagonal_offer_ribbon(c, x, y, w, h, offset=24, band_w=16, font_sz=7)

    # --------------------------------------------------------------------------
    # ZONA 2: CUERPO CENTRAL (IMAGEN CUADRADA A LA IZQ + TÍTULO Y DESC AMPLIADA)
    # --------------------------------------------------------------------------
    footer_bar_y = y + (12 if print_version else 68)
    body_h = header_bar_y - footer_bar_y  # 126 pt disponibles

    img_size = min(110, body_h - 14)
    img_x = x + 10
    img_y = footer_bar_y + (body_h - img_size) / 2

    fit_image(c, p["img"], img_x, img_y, img_size, img_size)

    product_target_url = p["url"] or WEB_URL
    link_url_new_tab(c, product_target_url, (img_x, img_y, img_x + img_size, img_y + img_size), print_version)

    tx = img_x + img_size + 12
    right_edge = x + w - 10

    # Título del producto (más grande y destacado)
    ty = header_bar_y - 17
    c.setFillColor(DARK)
    c.setFont("Helvetica-Bold", 12)
    nombre_lines = wrapped_lines(p["name"], 21, 3)
    for line in nombre_lines:
        c.drawString(tx, ty, line)
        ty -= 14

    title_h = len(nombre_lines) * 14
    link_url_new_tab(c, product_target_url, (tx, ty + 2, right_edge, ty + 2 + title_h), print_version)

    # Descripción ampliada con mayor tamaño de fuente y aprovechando el espacio
    ty -= 3
    c.setFillColor(HexColor("#4A6055"))
    c.setFont("Helvetica", 8.2)
    max_desc_lines = max(2, int((ty - (footer_bar_y + 6)) / 10.5))
    desc_lines = wrapped_lines(p["desc"], 28, max_desc_lines)
    for line in desc_lines:
        c.drawString(tx, ty, line)
        ty -= 10.5

    if print_version:
        return

    # Línea divisoria horizontal sobre el pie
    c.setStrokeColor(LINE)
    c.setLineWidth(0.8)
    c.line(x, footer_bar_y, x + w, footer_bar_y)

    # --------------------------------------------------------------------------
    # ZONA 3: PIE DE TARJETA (PRECIO A LA IZQ | BOTÓN VER EN LA WEB A LA DER)
    # --------------------------------------------------------------------------
    div_foot_x = x + 118
    c.setStrokeColor(LINE)
    c.setLineWidth(0.8)
    c.line(div_foot_x, y + 8, div_foot_x, footer_bar_y - 8)

    # Precio en el lado izquierdo
    if p.get("old_price"):
        old_str = money(p["old_price"])
        c.setFillColor(MID)
        c.setFont("Helvetica-Bold", 6.8)
        c.drawString(x + 12, y + 47, "ANTES")

        c.setFont("Helvetica-Bold", 8)
        c.setFillColor(HexColor("#85958D"))
        old_x = x + 38
        c.drawString(old_x, y + 47, old_str)

        old_w = c.stringWidth(old_str, "Helvetica-Bold", 8)
        c.setStrokeColor(RED)
        c.setLineWidth(1.1)
        c.line(old_x - 1, y + 49.5, old_x + old_w + 1, y + 49.5)

        if p["old_price"] > 0 and p.get("price") is not None:
            try:
                curr_n = float(p["price"])
                if curr_n < p["old_price"]:
                    desc_pct = int(round((1 - curr_n / p["old_price"]) * 100))
                    if desc_pct > 0:
                        pct_str = f"-{desc_pct}%"
                        pct_x = old_x + old_w + 5
                        pct_w = c.stringWidth(pct_str, "Helvetica-Bold", 6.5) + 6
                        c.setFillColor(HexColor("#FEE2E2"))
                        c.roundRect(pct_x, y + 44.5, pct_w, 10, 2, fill=1, stroke=0)
                        c.setFillColor(RED)
                        c.setFont("Helvetica-Bold", 6.5)
                        c.drawString(pct_x + 3, y + 47, pct_str)
            except (ValueError, TypeError):
                pass

        c.setFillColor(DARK)
        c.setFont("Helvetica-Bold", 19)
        c.drawString(x + 12, y + 20, money(p["price"]))
    else:
        c.setFillColor(MID)
        c.setFont("Helvetica-Bold", 7.2)
        c.drawString(x + 12, y + 47, "PRECIO")

        c.setFillColor(DARK)
        c.setFont("Helvetica-Bold", 20)
        c.drawString(x + 12, y + 20, money(p["price"]))

    # Botón único destacado: "Ver en la web >" en verde institucional ocupando todo su recuadro
    bx = div_foot_x + 8
    btn_w = x + w - 10 - bx
    btn_h = 36
    btn_y = y + (footer_bar_y - y - btn_h) / 2  # y + 16

    c.setFillColor(GREEN)
    c.roundRect(bx, btn_y, btn_w, btn_h, 6, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(bx + btn_w / 2, btn_y + 12.5, "Ver en la web >")
    link_url_new_tab(c, product_target_url, (bx, btn_y, bx + btn_w, btn_y + btn_h))


# ==============================================================================
# DISEÑO DE TARJETA HORIZONTAL (1 COLUMNA, 3 PRODUCTOS POR PÁGINA)
# ==============================================================================
def draw_card_horizontal(c, p, x, y, w, h, print_version=False):
    c.setFillColor(white)
    c.setStrokeColor(LINE)
    c.setLineWidth(0.8)
    c.roundRect(x, y, w, h, 10, fill=1, stroke=1)

    c.setFillColor(GREEN)
    c.roundRect(x, y, 7, h, 3.5, fill=1, stroke=0)

    img_x = x + 18
    img_y = y + 16
    img_w = 132
    img_h = h - 32
    fit_image(c, p["img"], img_x, img_y, img_w, img_h)

    product_target_url = p["url"] or WEB_URL
    link_url_new_tab(c, product_target_url, (img_x, img_y, img_x + img_w, img_y + img_h), print_version)

    div_x = img_x + img_w + 16
    c.setStrokeColor(LINE)
    c.line(div_x, y + 18, div_x, y + h - 18)

    tx = div_x + 18
    ty = y + h - 26

    c.setFillColor(GREEN)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(tx, ty, f"SKU {p['sku']}" if p["sku"] else "A PEDIDO")

    if p.get("offer") and not print_version:
        draw_diagonal_offer_ribbon(c, x, y, w, h, offset=28, band_w=20, font_sz=8)

    tags = []
    if p["presentation"]:
        tags.append(p["presentation"])
    if p["fragrances"]:
        tags.append(p["fragrances"])

    if tags:
        tag_text = " · ".join(tags)
        max_chars = 34 if p.get("offer") and not print_version else 50
        if len(tag_text) > max_chars:
            tag_text = tag_text[:max_chars - 3] + "..."
        c.setFillColor(LIGHT)
        max_tag_w = 165 if p.get("offer") and not print_version else 250
        tag_w = min(max_tag_w, c.stringWidth(tag_text, "Helvetica-Bold", 7) + 16)
        c.roundRect(tx + 60, ty - 5, tag_w, 16, 8, fill=1, stroke=0)
        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 7)
        c.drawString(tx + 68, ty, tag_text)

    ty -= 29
    c.setFillColor(DARK)
    c.setFont("Helvetica-Bold", 13.5)
    nombre_lines = wrapped_lines(p["name"], 38, 2)
    for line in nombre_lines:
        c.drawString(tx, ty, line)
        ty -= 15

    title_h = len(nombre_lines) * 16
    link_url_new_tab(c, product_target_url, (tx, ty + 10, tx + 240, ty + 10 + title_h), print_version)

    ty -= 3
    c.setFillColor(MID)
    c.setFont("Helvetica", 8.2)
    max_desc_lines = max(1, int((ty - (y + 18)) / 11) + 1) if print_version else 4
    for line in wrapped_lines(p["desc"], 58, max_desc_lines):
        c.drawString(tx, ty, line)
        ty -= 11

    if print_version:
        return

    if p.get("old_price"):
        old_str = money(p["old_price"])
        c.setFillColor(MID)
        c.setFont("Helvetica-Bold", 6.8)
        c.drawString(tx, y + 55, "ANTES")

        c.setFont("Helvetica-Bold", 9.5)
        c.setFillColor(HexColor("#85958D"))
        old_x = tx + 30
        c.drawString(old_x, y + 54, old_str)

        old_w = c.stringWidth(old_str, "Helvetica-Bold", 9.5)
        c.setStrokeColor(RED)
        c.setLineWidth(1.3)
        c.line(old_x - 2, y + 57.5, old_x + old_w + 2, y + 57.5)

        if p["old_price"] > 0 and p.get("price") is not None:
            try:
                curr_n = float(p["price"])
                if curr_n < p["old_price"]:
                    desc_pct = int(round((1 - curr_n / p["old_price"]) * 100))
                    if desc_pct > 0:
                        pct_str = f"-{desc_pct}%"
                        pct_x = old_x + old_w + 8
                        pct_w = c.stringWidth(pct_str, "Helvetica-Bold", 7) + 8
                        c.setFillColor(HexColor("#FEE2E2"))
                        c.roundRect(pct_x, y + 52, pct_w, 12, 3, fill=1, stroke=0)
                        c.setFillColor(RED)
                        c.setFont("Helvetica-Bold", 7)
                        c.drawString(pct_x + 4, y + 55, pct_str)
            except (ValueError, TypeError):
                pass

        c.setFillColor(RED if p.get("offer") else GREEN)
        c.setFont("Helvetica-Bold", 7)
        c.drawString(tx, y + 38, "OFERTA HOY" if p.get("offer") else "PRECIO AHORA")

        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 20)
        c.drawString(tx, y + 17, money(p["price"]))
    else:
        c.setFillColor(MID)
        c.setFont("Helvetica-Bold", 7.2)
        c.drawString(tx, y + 46, "PRECIO")

        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 20)
        c.drawString(tx, y + 21, money(p["price"]))

    btn_w = 145
    btn_h = 32
    btn_x = x + w - btn_w - 18
    btn_y = y + 18

    c.setFillColor(GREEN)
    c.roundRect(btn_x, btn_y, btn_w, btn_h, 6, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 11.5)
    c.drawCentredString(btn_x + btn_w / 2, btn_y + 11, "Ver en la web >")
    link_url_new_tab(c, product_target_url, (btn_x, btn_y, btn_x + btn_w, btn_y + btn_h))


def draw_custom_product(c, p, print_version=False):
    """Ficha independiente para productos a medida, sin SKU ni precio fijo."""
    x, y, w, h = 20, 48, W - 40, H - 156
    c.setFillColor(white)
    c.setStrokeColor(LINE)
    c.roundRect(x, y, w, h, 10, fill=1, stroke=1)

    c.setFillColor(LIGHT)
    c.roundRect(x + 20, y + h - 36, 147, 20, 6, fill=1, stroke=0)
    c.setFillColor(GREEN)
    c.setFont("Helvetica-Bold", 8.5)
    c.drawString(x + 31, y + h - 29, "FABRICACIÓN A PEDIDO")

    def paragraph(text, top, font_size=10.5, bold=False, color=DARK):
        style = ParagraphStyle(
            "personalizado", fontName="Helvetica-Bold" if bold else "Helvetica",
            fontSize=font_size, leading=font_size * 1.35, textColor=color,
        )
        block = Paragraph(escape(text).replace("\n", "<br/>"), style)
        _, height = block.wrap(w - 40, h)
        block.drawOn(c, x + 20, top - height)
        return top - height

    title_top = y + h - 52
    top = paragraph(p["name"], title_top, 23, bold=True)
    link_url_new_tab(c, p["url"] or WEB_URL, (x + 20, top, x + w - 20, title_top), print_version)
    top = paragraph("Tu logo, una bienvenida con identidad propia.", top - 6, 11, color=MID)

    image_h = 254
    image_y = top - image_h - 16
    fit_image(c, p["img"], x + 20, image_y, w - 40, image_h)
    link_url_new_tab(c, p["url"] or WEB_URL, (x + 20, image_y, x + w - 20, image_y + image_h), print_version)

    top = paragraph("Cotización según tu diseño", image_y - 17, 15, bold=True)
    paragraph(p["desc"], top - 10, 11)

    c.setFillColor(LIGHT)
    c.roundRect(x + 20, y + 77, w - 40, 44, 7, fill=1, stroke=0)
    label = "FABRICACIÓN A MEDIDA · SEÑA DEL 50% AL CONFIRMAR" if print_version else "PRECIO A COTIZAR · FABRICACIÓN CON SEÑA DEL 50%"
    paragraph(label, y + 105, 10, bold=True, color=GREEN)

    if print_version:
        paragraph("Consultá en el local para cotizar tu felpudo personalizado.", y + 47, 11, bold=True)
        return

    buttons = [("Ver en la web >", p["url"] or WEB_URL, 158)]
    if WHATSAPP_NUMERO:
        message = (
            "¡Hola! Quisiera cotizar un felpudo personalizado. "
            "Les paso las medidas, si lo quiero con borde o sin borde y mi logo vectorizado en Corel o SVG."
        )
        buttons.insert(0, ("Cotizar por WhatsApp >", f"https://wa.me/{WHATSAPP_NUMERO}?text={urllib.parse.quote(message)}", w - 218))
    bx = x + 20
    for label, url, bw in buttons:
        c.setFillColor(GREEN)
        c.roundRect(bx, y + 20, bw, 38, 6, fill=1, stroke=0)
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 11)
        c.drawCentredString(bx + bw / 2, y + 34, label)
        link_url_new_tab(c, url, (bx, y + 20, bx + bw, y + 58))
        bx += bw + 20


def product_batches(products, cards_per_page):
    """Mantiene el orden y reserva una página para cada producto personalizado."""
    batches = []
    pending = []
    for product in products:
        if product.get("custom"):
            if pending:
                batches.append(pending)
                pending = []
            batches.append([product])
        else:
            pending.append(product)
            if len(pending) == cards_per_page:
                batches.append(pending)
                pending = []
    if pending:
        batches.append(pending)
    return batches


def get_blended_cover_path():
    """Genera o recupera de caché la portada compuesta de portadin.png con camioneta.jpg difuminada en el fondo superior."""
    if not PORTADIN_IMG.exists():
        return None
    if not CAMIONETA_IMG.exists():
        return PORTADIN_IMG

    out_path = PORTADA_CACHE
    if (out_path.exists() and
        out_path.stat().st_mtime >= PORTADIN_IMG.stat().st_mtime and
        out_path.stat().st_mtime >= CAMIONETA_IMG.stat().st_mtime):
        return out_path

    try:
        port = Image.open(PORTADIN_IMG).convert("RGBA")
        pw, ph = port.size

        cam = Image.open(CAMIONETA_IMG).convert("RGBA")
        cw, ch = cam.size

        # Escalar camioneta al ancho exacto de portadin
        scale = pw / cw
        new_ch = int(ch * scale)
        cam_resized = cam.resize((pw, new_ch), Image.Resampling.LANCZOS)

        # Tomar la franja superior donde se aprecian las camionetas
        offset_y = 60
        target_h = 350
        cam_cropped = cam_resized.crop((0, offset_y, pw, offset_y + target_h))

        # Crear máscara de degradado vertical suave que se desvanece antes del logo
        vals = []
        for y in range(target_h):
            if y < 100:
                val = 230
            elif y < 320:
                p = (y - 100) / (320 - 100)
                val = int(230 * (1 - p) ** 1.8)
            else:
                val = 0
            vals.append(val)

        grad = Image.new("L", (1, target_h))
        grad.putdata(vals)
        mask = grad.resize((pw, target_h), Image.Resampling.NEAREST)

        cam_cropped.putalpha(mask)

        res = port.copy()
        res.alpha_composite(cam_cropped, (0, 0))
        out_path.parent.mkdir(parents=True, exist_ok=True)
        res.convert("RGB").save(out_path, quality=95)
        return out_path
    except Exception:
        return PORTADIN_IMG


def draw_cover(c, print_version=False):
    cover_img_path = get_blended_cover_path()
    if cover_img_path and cover_img_path.exists():
        # Portada fiel a portadin.png con camioneta.jpg en el fondo superior
        c.drawImage(str(cover_img_path), 0, 0, width=W, height=H)

        # Disclaimer inferior integrado discretamente con tipografía cuidada
        c.setFont("Helvetica-Bold", 9.5)
        c.setFillColor(HexColor("#D1FAE5"))
        if not print_version:
            c.drawCentredString(W / 2, 40, f"* {DISCLAIMER_PORTADA} *")
        return

    # Fallback si no está portadin.png (diseño institucional vectorial)
    c.setFillColor(GREEN)
    c.rect(0, 0, W, H, fill=1, stroke=0)

    c.setStrokeColor(HexColor("#1E8552"))
    c.setLineWidth(1.5)
    c.roundRect(18, 18, W - 36, H - 36, 14, fill=0, stroke=1)

    logo_size = 250
    logo_x = (W - logo_size) / 2
    logo_y = H - 340
    if LOGO.exists():
        c.drawImage(str(LOGO), logo_x, logo_y, width=logo_size, height=logo_size, mask="auto")

    title_y = logo_y - 42
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 30)
    c.drawCentredString(W / 2, title_y, TITULO_PORTADA)

    c.setFont("Helvetica-Bold", 13.5)
    c.setFillColor(HexColor("#EEF5F1"))
    c.drawCentredString(W / 2, title_y - 38, f"{NOMBRE_COMERCIO} · Mar del Plata")

    c.setFont("Helvetica-Bold", 9.5)
    c.setFillColor(HexColor("#FEF08A"))
    if not print_version:
        c.drawCentredString(W / 2, 55, f"* {DISCLAIMER_PORTADA} *")


def generate(print_version=False, output=None):
    output = Path(output) if output is not None else (SALIDA_IMPRESION if print_version else SALIDA)
    products = load_products()
    if not products:
        raise RuntimeError("No hay productos con Mostrar = SI.")

    header_h = 96
    footer_h = 26

    if COLUMNAS == 2:
        cards_per_page = 6
        card_w = 271
        card_h = 222
        col_gap = 12
        row_gap = 12
        start_x = 20
        top_y = H - header_h - 12
    else:
        cards_per_page = 3
        card_w = W - 52
        card_h = 200
        gap = 16
        top_y = H - header_h - 18

    batches = product_batches(products, cards_per_page)
    product_pages = len(batches)

    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        c = canvas.Canvas(str(output), pagesize=A4)
        c.setTitle(f"Catálogo {NOMBRE_COMERCIO}" + (" - Para imprimir" if print_version else ""))

        # Portada institucional inicial
        if INCLUIR_PORTADA:
            draw_cover(c, print_version)
            c.showPage()

        page_no = 1
        for batch in batches:
            draw_header(c, page_no, product_pages, print_version)

            if batch[0].get("custom"):
                draw_custom_product(c, batch[0], print_version)
            elif COLUMNAS == 2:
                for i, p in enumerate(batch):
                    row = i // 2
                    col = i % 2
                    x = start_x + col * (card_w + col_gap)
                    y = top_y - (row + 1) * card_h - row * row_gap
                    draw_card_grid(c, p, x, y, card_w, card_h, print_version)
            else:
                y_top = top_y
                for p in batch:
                    y = y_top - card_h
                    draw_card_horizontal(c, p, 26, y, card_w, card_h, print_version)
                    y_top = y - gap

            draw_footer(c, print_version)
            c.showPage()
            page_no += 1

        c.save()
    except PermissionError:
        raise PermissionError(
            f"No se pudo guardar '{output.name}'. El archivo está ABIERTO en otra aplicación "
            f"(Adobe Reader, navegador, etc.).\nPor favor ciérralo y vuelve a ejecutar."
        )

    size_kb = output.stat().st_size / 1024
    if size_kb >= 1024:
        size_str = f"{size_kb / 1024:.2f} MB"
    else:
        size_str = f"{size_kb:.0f} KB"

    total_pdf_pages = product_pages + (1 if INCLUIR_PORTADA else 0)
    print(f"Catálogo generado con éxito: {output} ({size_str}, {total_pdf_pages} pág{'s' if total_pdf_pages > 1 else ''})")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generador de catálogos de El Palacio de la Limpieza")
    parser.add_argument("--imprimir", action="store_true", help="Generar un PDF aparte, sin precios ni botones, para imprimir en el local")
    args = parser.parse_args()
    try:
        generate(print_version=args.imprimir)
    except PermissionError as pe:
        print("\n" + "=" * 65)
        print("AVISO: ARCHIVO PDF BLOQUEADO")
        print(pe)
        print("=" * 65 + "\n")
        sys.exit(1)
    except Exception as e:
        print(f"\nERROR: {e}\n")
        sys.exit(1)
