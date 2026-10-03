"""
Script para sincronizar automáticamente las imágenes de la carpeta 'imagenes/'
con la planilla 'catalogo_productos.xlsx'.

Si agregas fotos nuevas a la carpeta 'imagenes/' (nombradas por SKU, ej: 5003.png),
al ejecutar este script se agregan automáticamente las filas correspondientes
al Excel para que no tengas que escribir rutas ni códigos a mano.
"""
from pathlib import Path
from openpyxl import load_workbook
import sys

BASE = Path(__file__).resolve().parent
EXCEL = BASE / "catalogo_productos.xlsx"
IMG_DIR = BASE / "imagenes"

VALID_EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def sincronizar():
    if not EXCEL.exists():
        print(f"ERROR: No se encontró el archivo {EXCEL}")
        return

    if not IMG_DIR.exists():
        print(f"ERROR: No se encontró la carpeta {IMG_DIR}")
        return

    wb = load_workbook(EXCEL)
    ws = wb["Productos"]

    headers = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if v:
            headers[str(v).strip()] = col

    # Obtener SKUs existentes
    sku_col = headers.get("SKU")
    if not sku_col:
        print("ERROR: No se encontró la columna SKU en el Excel.")
        return

    skus_existentes = set()
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row, sku_col).value
        if val is not None and str(val).strip() != "":
            skus_existentes.add(str(val).strip())

    # Buscar imágenes en imagenes/
    imagenes_encontradas = []
    for f in IMG_DIR.iterdir():
        if f.is_file() and f.suffix.lower() in VALID_EXTS:
            sku_candidato = f.stem
            if sku_candidato not in skus_existentes:
                imagenes_encontradas.append((sku_candidato, f))

    if not imagenes_encontradas:
        print("¡Todo al día! No hay fotos nuevas sin registrar en el Excel.")
        return

    print(f"Se encontraron {len(imagenes_encontradas)} imágenes nuevas sin registrar:")
    for sku, f in imagenes_encontradas:
        print(f"  - SKU: {sku} ({f.name})")

    # Preguntar si desea ingresar datos interactivos o solo crear las filas
    print("\n¿Deseas ingresar nombre y precio ahora desde la consola? (s/n)")
    resp = input("> ").strip().lower()

    nuevas_filas = 0
    next_row = ws.max_row + 1

    # Asegurarnos de que next_row sea realmente la primera fila vacía
    while ws.cell(next_row - 1, sku_col).value is None and next_row > 2:
        next_row -= 1
    if ws.cell(next_row, sku_col).value is not None:
        next_row += 1

    for sku, f in imagenes_encontradas:
        rel_img = f"imagenes/{f.name}"
        nombre = f"Producto {sku}"
        desc = ""
        precio = 0
        presentacion = ""
        url_slug = sku

        if resp in ("s", "si", "sí", "y", "yes"):
            print(f"\n--- Cargando SKU {sku} ({f.name}) ---")
            nom_in = input("Nombre del producto: ").strip()
            if nom_in:
                nombre = nom_in

            prec_in = input("Precio (solo números): ").strip()
            if prec_in:
                try:
                    precio = float(prec_in.replace(".", "").replace(",", "."))
                except ValueError:
                    precio = prec_in

            desc_in = input("Descripción (opcional): ").strip()
            if desc_in:
                desc = desc_in

            pres_in = input("Presentación (ej. 5 litros) (opcional): ").strip()
            if pres_in:
                presentacion = pres_in

            url_in = input("Slug URL web (opcional, ej. detergente-5l): ").strip()
            if url_in:
                url_slug = url_in

        # Escribir en la hoja Excel
        def set_val(col_name, val):
            if col_name in headers:
                ws.cell(next_row, headers[col_name], val)

        set_val("SKU", sku)
        set_val("Producto", nombre)
        set_val("Descripción", desc)
        set_val("Precio", precio)
        set_val("Imagen", rel_img)
        set_val("Presentación", presentacion)
        set_val("Mostrar", "SI")
        set_val("URL", url_slug)

        next_row += 1
        nuevas_filas += 1

    wb.save(EXCEL)
    print(f"\n¡Éxito! Se agregaron {nuevas_filas} productos a '{EXCEL.name}'.")
    print("Ya puedes revisar el Excel o generar el catálogo directamente.")


if __name__ == "__main__":
    try:
        sincronizar()
    except KeyboardInterrupt:
        print("\nOperación cancelada.")
    except Exception as e:
        print(f"\nERROR: {e}")
        sys.exit(1)
